from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, datetime, timedelta
from io import BytesIO
from threading import Barrier

import pytest
import test_applications as m2
import test_m3 as m3
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import delete, func, select
from sqlalchemy import event as sql_event
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.applications import workflow
from app.applications.models import (
    Badge,
    BusinessAudit,
    CreatorApplication,
    LocalRoleAssignment,
    NotificationOutbox,
)
from app.config import settings
from app.creators import images
from app.creators.models import BannedChannel, CreatorProfile, ProfileImage
from app.events.models import Event
from app.helpers.models import HelperInvitation, HelperRegistration
from app.identity.roles import set_role
from app.main import app
from app.moderation import service as moderation
from app.notifications import client as notification_client
from app.notifications import worker
from app.registration.client import RegistrationStatus, RegistrationUnavailable
from app.staff import exports, service

application_engine = m2.application_engine
postgres_engine = m2.postgres_engine
browser = m2.browser
run = m3.run


class ProviderFake:
    def __init__(self, db=None, failure=None):
        self.db, self.failure = db, failure
        self.messages = []

    async def send(self, notification):
        if self.db is not None:
            assert not self.db.in_transaction()
        if self.failure:
            raise self.failure
        self.messages.append(notification)


class DownloadStore:
    def __init__(self):
        self.objects = {}

    def get(self, key):
        try:
            return self.objects[key]
        except KeyError:
            raise images.ImageStorageUnavailable from None


def complete_creator(db, user_id=1, store=None):
    app_id = m3.approve(db, user_id)
    store = store or DownloadStore()
    image = ProfileImage(
        event_id=1,
        object_key=f"profiles/test-{app_id}.png",
        state="ACTIVE",
        delete_after=datetime.now(UTC),
    )
    db.add(image)
    db.flush()
    profile = db.get(CreatorProfile, app_id)
    profile.channel_name, profile.image_id = f"Creator {user_id}", image.id
    db.get(CreatorApplication, app_id).staff_notes = "Private staff context"
    store.objects[image.object_key] = images.normalize_png(m3.png())
    db.commit()
    return app_id, store


def workbook(data):
    return load_workbook(BytesIO(data), data_only=False)


@pytest.mark.parametrize("kind", list(worker.MESSAGES))
def test_dispatch_all_kinds_and_no_completed_replay(application_engine, kind):
    with Session(application_engine, expire_on_commit=False) as db:
        app_id = m3.approve(db)
        row = db.scalar(select(NotificationOutbox))
        row.notification_type = kind
        db.commit()
        provider = ProviderFake(db)
        assert run(worker.dispatch(db, provider)) == 1
        assert run(worker.dispatch(db, provider)) == 0
        assert len(provider.messages) == 1
        message = provider.messages[0]
        assert message.kind == kind and message.application_id == app_id
        assert message.category == "Operational"
        db.refresh(row)
        assert row.state == "SENT" and row.attempts == 1 and row.sent_at
        assert row.claim_id is None


def test_notification_workflow_events(application_engine):
    with Session(application_engine, expire_on_commit=False) as db:
        app_id = m3.approve(db)
        helper_id = m3.register(db, app_id)
        m3.decide(db, helper_id)
        m3.decide(db, helper_id, "DECLINED")
        m3.decide(db, helper_id)
        m3.withdraw(db, app_id)
        m3.withdraw(
            db, app_id, actor=2, administrative=True, restore=True, reason="Correction"
        )
        record = db.get(CreatorApplication, app_id)
        run(
            workflow.review(
                db,
                2,
                app_id,
                record.version,
                "ON_REVIEW",
                "Correction",
                "",
                False,
                m2.RegistrationFake(),
                m2.IdentityFake(),
            )
        )
        provider = ProviderFake(db)
        run(worker.dispatch(db, provider))
        kinds = [n.kind for n in provider.messages]
        assert "ON_REVIEW" not in kinds
        assert "APPROVAL_REVOKED" in kinds
        assert kinds.count("CREATOR_INACTIVE") == 2
        assert {
            "HELPER_REQUESTED",
            "HELPER_CONFIRMED",
            "HELPER_DECLINED",
            "CREATOR_RESTORED",
        } <= set(kinds)
        assert (
            next(
                n for n in provider.messages if n.kind == "HELPER_REQUESTED"
            ).recipient_id
            == 1
        )
        assert all(
            n.recipient_id == 3
            for n in provider.messages
            if n.kind in ("HELPER_CONFIRMED", "HELPER_DECLINED", "CREATOR_INACTIVE")
        )


@pytest.mark.parametrize("target", ["NOT_APPROVED", "NOT_ACCEPTED"])
def test_rejection_intents_deliver(application_engine, target):
    with Session(application_engine, expire_on_commit=False) as db:
        app_id = run(
            workflow.submit(
                db, 1, m2.input_data(), m2.RegistrationFake(), m2.IdentityFake()
            )
        )
        if target == "NOT_APPROVED":
            run(
                workflow.review(
                    db,
                    2,
                    app_id,
                    1,
                    "ON_REVIEW",
                    "",
                    "",
                    False,
                    m2.RegistrationFake(),
                    m2.IdentityFake(),
                )
            )
            assert db.scalar(select(func.count()).select_from(NotificationOutbox)) == 0
        record = db.get(CreatorApplication, app_id)
        run(
            workflow.review(
                db,
                2,
                app_id,
                record.version,
                target,
                "Visible reason",
                "",
                False,
                m2.RegistrationFake(),
                m2.IdentityFake(),
            )
        )
        provider = ProviderFake(db)
        run(worker.dispatch(db, provider))
        assert [n.kind for n in provider.messages] == [target]


def test_notification_retry_limit_retry_after_sanitization(application_engine):
    with Session(application_engine, expire_on_commit=False) as db:
        app_id = m3.approve(db)
        provider = ProviderFake(
            db, notification_client.DeliveryFailure(retry_after=600)
        )
        now = datetime.now(UTC)
        for attempt in range(1, 6):
            assert run(worker.dispatch(db, provider, now=now)) == 1
            row = db.scalar(
                select(NotificationOutbox).execution_options(populate_existing=True)
            )
            assert row.attempts == attempt
            assert db.get(CreatorApplication, app_id).status == "APPROVED"
            assert db.scalar(select(Badge)).badge_number == 1
            if attempt < 5:
                assert workflow.utc(row.next_attempt_at) >= now + timedelta(seconds=600)
                assert (
                    run(worker.dispatch(db, provider, now=now + timedelta(seconds=1)))
                    == 0
                )
                now = workflow.utc(row.next_attempt_at)
        assert row.state == "FAILED"
        assert run(worker.dispatch(db, provider, now=now + timedelta(days=1))) == 0
        assert db.scalar(
            select(BusinessAudit).where(BusinessAudit.action == "notification_failed")
        )


def test_success_after_retry_and_expired_lease(application_engine):
    with Session(application_engine, expire_on_commit=False) as db:
        m3.approve(db)
        now = datetime.now(UTC)
        run(
            worker.dispatch(
                db, ProviderFake(db, RuntimeError("secret-do-not-store")), now=now
            )
        )
        row = db.scalar(
            select(NotificationOutbox).execution_options(populate_existing=True)
        )
        assert row.last_error == "DELIVERY_UNCONFIRMED"
        db.rollback()
        claimed = worker.claim(db, now + timedelta(minutes=2))
        assert claimed
        assert worker.claim(db, now + timedelta(minutes=3)) is None
        provider = ProviderFake(db)
        run(worker.dispatch(db, provider, now=now + timedelta(minutes=8)))
        db.refresh(row)
        assert row.state == "SENT" and row.attempts == 3
        assert len(provider.messages) == 1


def test_unavailable_provider_is_explicit(application_engine, monkeypatch):
    monkeypatch.setattr(settings, "notification_provider_factory", None)
    provider = notification_client.get_notification_provider()
    assert isinstance(provider, notification_client.UnavailableNotificationProvider)
    with Session(application_engine, expire_on_commit=False) as db:
        m3.approve(db)
        run(worker.dispatch(db, provider))
        row = db.scalar(select(NotificationOutbox))
        assert row.last_error == "PROVIDER_UNAVAILABLE" and row.state == "PENDING"
    monkeypatch.setattr(
        settings, "notification_provider_factory", "test_m4:ProviderFake"
    )
    assert isinstance(notification_client.get_notification_provider(), ProviderFake)
    monkeypatch.setattr(
        settings, "notification_provider_factory", "missing_module:secret"
    )
    with pytest.raises(RuntimeError, match="configuration is unavailable"):
        notification_client.get_notification_provider()


def test_expired_final_claim_fails_without_sending(application_engine):
    with Session(application_engine, expire_on_commit=False) as db:
        m3.approve(db)
        row = db.scalar(select(NotificationOutbox))
        row.state, row.attempts, row.claimed_until = (
            "SENDING",
            5,
            datetime.now(UTC) - timedelta(minutes=1),
        )
        db.commit()
        provider = ProviderFake(db)
        run(worker.dispatch(db, provider))
        db.refresh(row)
        assert row.state == "FAILED" and not provider.messages


def test_print_rows_operations_and_no_secrets(application_engine):
    with Session(application_engine, expire_on_commit=False) as db:
        app_id, store = complete_creator(db)
        helper_id = m3.register(db, app_id)
        m3.decide(db, helper_id)
        invitation = db.scalar(select(HelperInvitation))
        digest = invitation.token_digest
        registration = m2.RegistrationFake()
        data = run(exports.export(db, 2, 1, registration, store))
        book = workbook(data)
        rows = list(book["Print"].values)
        assert rows[0] == exports.PRINT_COLUMNS
        assert rows[1] == (
            1,
            "Creator 1",
            "Twitch",
            "https://www.twitch.tv/creator1",
            f"creator1-{app_id}.png",
        )
        assert rows[2] == (2, *rows[1][1:])
        assert registration.calls == 2
        assert {
            "Event",
            "Applications",
            "Channels",
            "Videos",
            "Profiles",
            "Images",
            "Invitations",
            "Helpers",
            "Badges",
            "Roles",
            "Notifications",
            "Audit",
            "Ban warnings",
            "Counters",
        } <= set(book.sheetnames)
        values = str([list(sheet.values) for sheet in book])
        assert "Private staff context" in values and "print_exported" in values
        assert "attendee@example.test" in values and "REG-123" in values
        assert (
            digest not in values
            and "token_digest" not in values
            and "claim_id" not in values
        )
        assert (
            "profiles/test-" in values
        )  # Private object references aid recovery; never public access.


@pytest.mark.parametrize(
    "failure",
    [
        "INELIGIBLE",
        "UNAVAILABLE",
        "UNKNOWN",
        "missing_name",
        "missing_image",
        "missing_object",
        "bad_image",
        "missing_badge",
    ],
)
def test_print_validation_blocks_without_workflow_changes(application_engine, failure):
    with Session(application_engine, expire_on_commit=False) as db:
        app_id, store = complete_creator(db)
        helper_id = m3.register(db, app_id)
        m3.decide(db, helper_id)
        registration = m2.RegistrationFake()
        if failure == "INELIGIBLE":
            registration.status = RegistrationStatus.INELIGIBLE
        elif failure == "UNAVAILABLE":
            registration.failure = RegistrationUnavailable()
        elif failure == "UNKNOWN":
            registration.status = RegistrationStatus.UNKNOWN
        elif failure == "missing_name":
            db.get(CreatorProfile, app_id).channel_name = ""
        elif failure == "missing_image":
            db.get(CreatorProfile, app_id).image_id = None
        elif failure == "missing_object":
            store.objects.clear()
        elif failure == "bad_image":
            store.objects = {key: b"invalid" for key in store.objects}
        else:
            db.execute(delete(Badge).where(Badge.application_id == app_id))
        db.commit()
        with pytest.raises(HTTPException) as error:
            run(exports.export(db, 2, 1, registration, store))
        assert error.value.status_code == 409
        problems = error.value.detail["problems"]
        assert problems and problems[0]["application_id"] == app_id
        assert db.get(CreatorApplication, app_id).status == "APPROVED"
        assert db.get(HelperRegistration, helper_id).status == "CONFIRMED"
        fallback = workbook(
            run(exports.export(db, 2, 1, registration, store, print_ready=False))
        )
        assert (
            "Print" not in fallback.sheetnames
            and "Export status" in fallback.sheetnames
        )
        assert len(list(fallback["Applications"].values)) == 2


def test_inactive_print_exclusion_preserves_operations(application_engine):
    with Session(application_engine, expire_on_commit=False) as db:
        app_id, store = complete_creator(db)
        helper_id = m3.register(db, app_id)
        m3.decide(db, helper_id)
        m3.withdraw(db, helper_id, actor=3, helper=True)
        result = workbook(run(exports.export(db, 2, 1, m2.RegistrationFake(), store)))
        assert result["Print"].max_row == 2
        assert result["Helpers"].max_row == 2 and result["Badges"].max_row == 3
        m3.withdraw(db, app_id)
        result = workbook(run(exports.export(db, 2, 1, m2.RegistrationFake(), store)))
        assert result["Print"].max_row == 1
        assert result["Applications"].max_row == 2


def test_export_revalidates_changes_and_admin_revocation(application_engine):
    with Session(application_engine, expire_on_commit=False) as db:
        app_id, store = complete_creator(db)
        registration = m2.RegistrationFake()

        def change():
            assert not db.in_transaction()
            with Session(application_engine) as other:
                other.get(CreatorProfile, app_id).channel_name = "Changed during lookup"
                other.commit()

        registration.hook = change
        with pytest.raises(HTTPException, match="changed during verification"):
            run(exports.export(db, 2, 1, registration, store))

        def revoke():
            with Session(application_engine) as other:
                other.execute(
                    delete(LocalRoleAssignment).where(LocalRoleAssignment.user_id == 2)
                )
                other.commit()

        registration.hook = revoke
        with pytest.raises(HTTPException) as error:
            run(exports.export(db, 2, 1, registration, store))
        assert error.value.status_code == 403


@pytest.mark.parametrize(
    "value",
    ["=1+1", "+cmd", "-1+cmd", "@SUM(1)", "\t=1+1", "\n=1+1", "https://example.test"],
)
def test_every_string_is_literal(value):
    book = workbook(
        exports.workbook_bytes(
            {"Operations": [{"value": value, "nested": {"x": value}}]}
        )
    )
    cell = book["Operations"]["A2"]
    assert cell.value == value and cell.data_type == "s" and cell.hyperlink is None
    assert all(c.data_type != "f" for sheet in book for row in sheet for c in row)


def test_picture_download_safe_primary_filename_and_failure(application_engine):
    with Session(application_engine, expire_on_commit=False) as db:
        app_id, store = complete_creator(db)
        m3.withdraw(db, app_id)
        name, data = exports.download_picture(db, 2, app_id, store)
        assert name == f"creator1-{app_id}.png" and data.startswith(b"\x89PNG")
        assert db.scalar(
            select(BusinessAudit).where(BusinessAudit.action == "picture_downloaded")
        )
        for user in (1, 3):
            with pytest.raises(HTTPException) as error:
                exports.download_picture(db, user, app_id, store)
            assert error.value.status_code == 403
        store.objects.clear()
        with pytest.raises(HTTPException) as error:
            exports.download_picture(db, 2, app_id, store)
        assert error.value.status_code == 503 and "profiles/" not in str(
            error.value.detail
        )
    for account in ("@handle", "../../CON", 'a\r\n"/x', "☃", "same@social.example"):
        name = exports.picture_filename(
            type("Channel", (), {"normalized_account": account})(), 12
        )
        assert name.endswith("-12.png") and not any(c in name for c in '/\\\r\n"@')


def test_pickup_lookup_multiple_badges_and_correction(application_engine):
    with Session(application_engine, expire_on_commit=False) as db:
        app_id = m3.approve(db)
        helper_id = m3.register(db, app_id)
        m3.decide(db, helper_id)
        rows = service.lookup(db, 3, 1, "REG-123")
        assert [r["kind"] for r in rows] == ["Creator", "Helper"]
        assert not any("email" in r or "nickname" in r for r in rows)
        for row in rows:
            service.pickup(db, 3, 1, row["id"])
            service.pickup(db, 3, 1, row["id"])
        assert (
            db.scalar(
                select(func.count())
                .select_from(BusinessAudit)
                .where(BusinessAudit.action == "picked_up")
            )
            == 2
        )
        badges = list(db.scalars(select(Badge)))
        assert all(b.picked_up_by == 3 and b.picked_up_at for b in badges)
        with pytest.raises(HTTPException) as error:
            service.pickup(db, 3, 1, rows[0]["id"], undo=True, reason="wrong")
        assert error.value.status_code == 403
        service.pickup(db, 2, 1, rows[0]["id"], undo=True, reason="Wrong badge")
        assert db.get(Badge, rows[0]["id"]).picked_up_at is None
        assert db.get(Badge, rows[1]["id"]).picked_up_at
        assert db.scalar(
            select(BusinessAudit).where(BusinessAudit.action == "pickup_corrected")
        )


def test_pickup_inactivity_roles_event_scope_and_revoke(application_engine):
    with Session(application_engine, expire_on_commit=False) as db:
        app_id = m3.approve(db)
        helper_id = m3.register(db, app_id)
        m3.decide(db, helper_id)
        rows = service.lookup(db, 3, 1, "REG-123")
        m3.withdraw(db, app_id)
        for row in rows:
            with pytest.raises(HTTPException) as error:
                service.pickup(db, 3, 1, row["id"])
            assert error.value.status_code == 409
        for user, event_id in ((1, 1), (3, 2)):
            with pytest.raises(HTTPException) as error:
                service.lookup(db, user, event_id, "REG-123")
            assert error.value.status_code == 403
        db.rollback()
        with db.begin():
            set_role(
                db,
                settings.oidc_issuer_url,
                "person-3",
                False,
                "Shift ended",
                event_id=1,
            )
        with pytest.raises(HTTPException):
            service.lookup(db, 3, 1, "REG-123")
        db.rollback()
        with db.begin():
            set_role(
                db, settings.oidc_issuer_url, "person-3", True, "New shift", event_id=1
            )
        assert len(service.lookup(db, 3, 1, "REG-123")) == 2


def test_ban_management_normalization_and_warning_only(application_engine):
    with Session(application_engine, expire_on_commit=False) as db:
        ban_id = moderation.save(
            db,
            2,
            platform="Twitch",
            account="https://www.twitch.tv/CREATOR1",
            reason="Private moderation reason",
        )
        app_id = m3.approve(db)
        assert db.get(CreatorApplication, app_id).status == "APPROVED"
        assert [b.id for b in moderation.warnings(db, app_id)] == [ban_id]
        moderation.save(
            db,
            2,
            ban_id=ban_id,
            platform="Twitch",
            account="@Creator1",
            reason="Updated reason",
            active=False,
        )
        assert not moderation.warnings(db, app_id)
        moderation.save(
            db,
            2,
            ban_id=ban_id,
            platform="Twitch",
            account="@Creator1",
            reason="Updated reason",
            active=True,
        )
        assert moderation.warnings(db, app_id)
        record = db.get(CreatorApplication, app_id)
        m3.service.save_profile(
            db,
            1,
            app_id,
            record.version,
            "Name",
            m2.input_data("creator1extra").channels,
        )
        assert not moderation.warnings(db, app_id)
        for platform, account in (
            ("X", "creator1"),
            ("Mastodon", "https://one.example/@same"),
            ("Mastodon", "https://two.example/@same"),
        ):
            moderation.save(
                db, 2, platform=platform, account=account, reason="Other account"
            )
        for user in (1, 3):
            with pytest.raises(HTTPException) as error:
                moderation.save(
                    db, user, platform="Twitch", account="x", reason="Denied"
                )
            assert error.value.status_code == 403
        with pytest.raises(HTTPException) as error:
            moderation.save(
                db, 2, platform="Twitch", account="creator1", reason="Duplicate"
            )
        assert error.value.status_code == 409
        moderation.delete_ban(db, 2, ban_id, "Lifted permanently")
        assert db.get(BannedChannel, ban_id) is None
        assert db.scalar(
            select(BusinessAudit).where(BusinessAudit.action == "ban_deleted")
        )


def test_ban_survives_event_and_actor_deletion(application_engine):
    from app.identity.models import ExternalIdentity, LocalUser

    with Session(application_engine, expire_on_commit=False) as db:
        ban_id = moderation.save(
            db, 2, platform="Twitch", account="retained", reason="Retained ban"
        )
        db.execute(delete(Event))
        db.execute(delete(ExternalIdentity))
        db.execute(delete(LocalUser))
        db.commit()
        assert db.get(BannedChannel, ban_id).private_reason == "Retained ban"
        assert not list(db.scalars(select(BusinessAudit)))


def test_notification_constraints(application_engine):
    with Session(application_engine, expire_on_commit=False) as db:
        m3.approve(db)
        row = db.scalar(select(NotificationOutbox))
        for field, value in (("attempts", 6), ("state", "anything")):
            setattr(row, field, value)
            with pytest.raises(IntegrityError):
                db.commit()
            db.rollback()


def test_admin_routes_export_download_warnings_and_private_staff_access(browser):
    client, _, engine = browser
    with Session(engine, expire_on_commit=False) as db:
        app_id, store = complete_creator(db)
        moderation.save(
            db, 2, platform="Twitch", account="creator1", reason="Private ban detail"
        )
    app.dependency_overrides[images.get_image_store] = lambda: store
    m2.login(client, 2)
    assert "Private ban detail" in client.get(f"/admin/applications/{app_id}").text
    response = client.get(f"/admin/pictures/{app_id}")
    assert (
        response.status_code == 200 and "no-store" in response.headers["cache-control"]
    )
    assert (
        response.headers["content-disposition"]
        == f'attachment; filename="creator1-{app_id}.png"'
    )
    token = m2.csrf(client, "/admin/events/1/operations")
    response = client.post(
        "/admin/events/1/export", data={"csrf_token": token, "mode": "print"}
    )
    assert (
        response.status_code == 200 and "no-store" in response.headers["cache-control"]
    )
    assert workbook(response.content)["Print"].max_row == 2
    assert (
        client.post("/admin/events/1/export", data={"mode": "print"}).status_code == 403
    )
    for user in (1, 3):
        m2.login(client, user)
        for path in (
            "/admin/events/1/operations",
            "/admin/bans",
            "/admin/notifications",
            f"/admin/pictures/{app_id}",
            f"/admin/applications/{app_id}",
        ):
            assert client.get(path).status_code == 403
        assert (
            client.post("/admin/events/1/export", data={"mode": "print"}).status_code
            == 403
        )
    m2.login(client, 1)
    assert "Private ban detail" not in client.get("/applications").text
    m2.login(client, 3)
    token = m2.csrf(client, "/staff/events/1/badges")
    response = client.post(
        "/staff/events/1/badges", data={"csrf_token": token, "reg_id": "REG-123"}
    )
    assert response.status_code == 200 and "Badge 1" in response.text
    with Session(engine) as db:
        badge_id = db.scalar(select(Badge.id))
    response = client.post(
        f"/staff/events/1/badges/{badge_id}",
        data={"csrf_token": token, "action": "pickup"},
        follow_redirects=False,
    )
    assert response.status_code == 303
    with Session(engine) as db:
        db.execute(delete(LocalRoleAssignment).where(LocalRoleAssignment.user_id == 3))
        db.commit()
    assert client.get("/staff/events/1/badges").status_code == 403


def test_postgres_concurrent_dispatch_and_pickup(postgres_engine, application_engine):
    if application_engine.dialect.name != "postgresql":
        pytest.skip("PostgreSQL locking test")
    with Session(application_engine, expire_on_commit=False) as db:
        m3.approve(db)
        badge_id = db.scalar(select(Badge.id))
    barrier = Barrier(2)
    provider = ProviderFake()

    def deliver():
        with Session(application_engine) as db:
            barrier.wait()
            return run(worker.dispatch(db, provider))

    with ThreadPoolExecutor(2) as pool:
        counts = list(pool.map(lambda _: deliver(), range(2)))
    assert sum(counts) == 1 and len(provider.messages) == 1
    barrier = Barrier(2)

    def collect():
        with Session(application_engine) as db:
            barrier.wait()
            service.pickup(db, 3, 1, badge_id)

    with ThreadPoolExecutor(2) as pool:
        list(pool.map(lambda _: collect(), range(2)))
    with Session(application_engine) as db:
        assert (
            db.scalar(
                select(func.count())
                .select_from(BusinessAudit)
                .where(BusinessAudit.action == "picked_up")
            )
            == 1
        )


def test_pickup_audit_failure_rolls_back(application_engine):
    with Session(application_engine, expire_on_commit=False) as db:
        m3.approve(db)
        badge_id = db.scalar(select(Badge.id))

        def fail_audit(session, *_):
            if any(
                isinstance(row, BusinessAudit) and row.action == "picked_up"
                for row in session.new
            ):
                raise RuntimeError("Audit unavailable")

        sql_event.listen(db, "before_flush", fail_audit)
        with pytest.raises(RuntimeError, match="Audit unavailable"):
            service.pickup(db, 3, 1, badge_id)
        sql_event.remove(db, "before_flush", fail_audit)
        assert db.get(Badge, badge_id).picked_up_at is None
        assert db.get(Badge, badge_id).picked_up_by is None


def test_actual_same_person_creator_and_helper_lookup(application_engine):
    with Session(application_engine, expire_on_commit=False) as db:
        own_id = m3.approve(db, 1)
        other_id = m3.approve(db, 2)
        helper_id = m3.register(db, other_id, user=1, actor=2)
        m3.decide(db, helper_id, actor=2)
        db.get(CreatorApplication, other_id).reg_id = "OTHER"
        db.commit()
        rows = service.lookup(db, 3, 1, "REG-123")
        assert [(r["kind"], r["application_id"]) for r in rows] == [
            ("Creator", own_id),
            ("Helper", other_id),
        ]
        assert len(service.lookup(db, 3, 1, "OTHER")) == 1
        assert not service.lookup(db, 3, 1, "missing")


@pytest.mark.parametrize(
    "platform,account,match",
    [
        ("Twitch", "same", False),
        ("Mastodon", "https://one.example/@same", True),
        ("Mastodon", "https://two.example/@same", False),
        ("Mastodon", "https://one.example/@same_extra", False),
    ],
)
def test_federated_ban_matching(application_engine, platform, account, match):
    with Session(application_engine, expire_on_commit=False) as db:
        app_id = m3.approve(db)
        moderation.save(
            db,
            2,
            platform="Mastodon",
            account="https://ONE.example/@SAME",
            reason="Specific account",
        )
        record = db.get(CreatorApplication, app_id)
        m3.service.save_profile(
            db,
            1,
            app_id,
            record.version,
            "Name",
            (m2.normalize_channel(platform, account, True),),
        )
        assert bool(moderation.warnings(db, app_id)) == match


def test_nonapproved_and_declined_print_exclusion(application_engine):
    with Session(application_engine, expire_on_commit=False) as db:
        app_id, store = complete_creator(db)
        helper_id = m3.register(db, app_id)
        assert (
            workbook(run(exports.export(db, 2, 1, m2.RegistrationFake(), store)))[
                "Print"
            ].max_row
            == 2
        )
        m3.decide(db, helper_id)
        m3.decide(db, helper_id, "DECLINED")
        result = workbook(run(exports.export(db, 2, 1, m2.RegistrationFake(), store)))
        assert result["Print"].max_row == 2 and result["Badges"].max_row == 3
        record = db.get(CreatorApplication, app_id)
        run(
            workflow.review(
                db,
                2,
                app_id,
                record.version,
                "ON_REVIEW",
                "Approval correction",
                "",
                False,
                m2.RegistrationFake(),
                m2.IdentityFake(),
            )
        )
        assert (
            workbook(run(exports.export(db, 2, 1, m2.RegistrationFake(), store)))[
                "Print"
            ].max_row
            == 1
        )


def test_secondary_channel_never_drives_print_row(application_engine):
    with Session(application_engine, expire_on_commit=False) as db:
        app_id, store = complete_creator(db)
        record = db.get(CreatorApplication, app_id)
        m3.service.save_profile(
            db,
            1,
            app_id,
            record.version,
            "Single display name",
            (
                m2.normalize_channel("X", "secondary", False),
                m2.normalize_channel("Mastodon", "https://social.example/@main", True),
            ),
        )
        result = workbook(run(exports.export(db, 2, 1, m2.RegistrationFake(), store)))
        row = list(result["Print"].values)[1]
        assert row[1:] == (
            "Single display name",
            "Mastodon",
            "https://social.example/@main",
            f"main_social.example-{app_id}.png",
        )
        assert exports.download_picture(db, 2, app_id, store)[0] == row[4]


def test_ban_forms_failed_notification_page_and_export_problem_ui(browser):
    client, registration, engine = browser
    with Session(engine, expire_on_commit=False) as db:
        app_id, store = complete_creator(db)
        run(
            worker.dispatch(
                db, ProviderFake(db, notification_client.ProviderUnavailable())
            )
        )
    app.dependency_overrides[images.get_image_store] = lambda: store
    m2.login(client, 2)
    token = m2.csrf(client, "/admin/bans")
    response = client.post(
        "/admin/bans",
        data={
            "csrf_token": token,
            "platform": "Twitch",
            "account": "creator1",
            "reason": "<script>private</script>",
            "active": "yes",
        },
    )
    assert (
        response.status_code == 200
        and "&lt;script&gt;private&lt;/script&gt;" in response.text
    )
    with Session(engine) as db:
        ban_id = db.scalar(select(BannedChannel.id))
    assert (
        client.post(
            "/admin/bans",
            data={
                "csrf_token": token,
                "ban_id": ban_id,
                "platform": "Twitch",
                "account": "creator1",
                "reason": "Retained",
                "action": "save",
            },
        ).status_code
        == 200
    )
    assert (
        client.post(
            "/admin/bans",
            data={
                "csrf_token": token,
                "ban_id": ban_id,
                "action": "delete",
                "reason": "Removed",
            },
        ).status_code
        == 200
    )
    assert "PROVIDER_UNAVAILABLE" in client.get("/admin/notifications").text
    registration.failure = RegistrationUnavailable()
    response = client.post(
        "/admin/events/1/export",
        data={
            "csrf_token": m2.csrf(client, "/admin/events/1/operations"),
            "mode": "print",
        },
    )
    assert (
        response.status_code == 409
        and "UNAVAILABLE" in response.text
        and "Badge 1" in response.text
    )
    assert "no-store" in response.headers["cache-control"]
    assert client.get(f"/creators/{app_id}?administrative=true").status_code == 200


def test_staff_cannot_mutate_admin_features_or_undo(browser):
    client, _, engine = browser
    with Session(engine, expire_on_commit=False) as db:
        app_id = m3.approve(db)
        helper_id = m3.register(db, app_id)
        m3.decide(db, helper_id)
        badge_id = db.scalar(select(Badge.id).where(Badge.application_id == app_id))
        version = db.get(CreatorApplication, app_id).version
    m2.login(client, 3)
    token = m2.csrf(client, "/staff/events/1/badges")
    for path in (
        "/admin/bans",
        "/admin/events/1/export",
        f"/admin/applications/{app_id}",
    ):
        assert client.post(path, data={"csrf_token": token}).status_code == 403
    assert (
        client.post(
            f"/creators/{app_id}/profile",
            data={
                "csrf_token": token,
                "channel_name": "stolen",
                "version": version,
                "platform": "Twitch",
                "account": "stolen",
                "primary": "0",
                "administrative": "yes",
                "reason": "denied",
            },
        ).status_code
        == 403
    )
    assert client.post(
        f"/helpers/{helper_id}/decision",
        data={"csrf_token": token, "version": 2, "status": "DECLINED"},
    ).status_code in (403, 404)
    assert (
        client.post(
            f"/staff/events/1/badges/{badge_id}",
            data={"csrf_token": token, "action": "undo", "reason": "Denied"},
        ).status_code
        == 403
    )
    assert (
        client.post(
            f"/staff/events/2/badges/{badge_id}",
            data={"csrf_token": token, "action": "pickup"},
        ).status_code
        == 403
    )
