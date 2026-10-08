import asyncio
import os
import subprocess
from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, datetime, timedelta
from threading import Barrier

import pytest
import test_applications as m2
import test_auth as auth_tests
import test_m3 as m3
from fastapi import HTTPException
from sqlalchemy import delete, func, select
from sqlalchemy.exc import OperationalError
from sqlalchemy.orm import Session

from app.applications import workflow
from app.applications.models import (
    Badge,
    BusinessAudit,
    CreatorApplication,
    CreatorChannel,
    LocalRoleAssignment,
    NotificationOutbox,
)
from app.creators import images
from app.creators.models import BannedChannel, CreatorProfile, ProfileImage
from app.database import Base, get_db
from app.events.cleanup import cleanup_due_events, cleanup_event
from app.events.models import Event
from app.gallery import service
from app.helpers.models import HelperRegistration
from app.identity.models import ExternalIdentity, LocalUser
from app.main import app
from app.notifications.worker import claim

application_engine = m2.application_engine
postgres_engine = m2.postgres_engine
browser = m2.browser
key = auth_tests.key
provider = auth_tests.provider


def published(db, user=1, name="Creator"):
    key = m3.approve(db, user)
    profile = db.get(CreatorProfile, key)
    profile.channel_name = name
    image = ProfileImage(
        event_id=1,
        object_key=f"private/{user}.png",
        state="ACTIVE",
        delete_after=datetime.now(UTC) + timedelta(days=60),
    )
    db.add(image)
    db.flush()
    profile.image_id = image.id
    db.commit()
    return key, profile.public_id, image.object_key


@pytest.fixture
def gallery_setup(browser):
    client, _, engine = browser
    with Session(engine, expire_on_commit=False) as db:
        key, public_id, _ = published(db, name="Zulu <creator>")
        published(db, 3, "Alpha")
        db.add(
            CreatorChannel(
                application_id=key,
                platform="Mastodon",
                original_representation="https://social.example/@extra",
                normalized_account="extra@social.example",
                canonical_url="https://social.example/@extra",
                is_primary=False,
            )
        )
        db.commit()

    class Store:
        failure = False
        hook = None

        def get(self, key):
            if self.failure:
                raise images.ImageStorageUnavailable
            if self.hook:
                self.hook()
            return m3.png()

    store = Store()
    app.dependency_overrides[images.get_image_store] = lambda: store
    client.cookies.clear()
    yield client, engine, key, public_id, store


def test_public_contract_order_html_and_conditional_requests(gallery_setup):
    client, _, key, public_id, _ = gallery_setup
    response = client.get("/api/v1/events/2028/creators")
    assert response.status_code == 200
    rows = response.json()
    assert [r["name"] for r in rows] == ["Alpha", "Zulu <creator>"]
    assert set(rows[1]) == {"id", "name", "channels", "image_url"}
    assert rows[1]["id"] == public_id and public_id != str(key)
    assert [c["primary"] for c in rows[1]["channels"]] == [True, False]
    assert set(rows[1]["channels"][0]) == {"platform", "account", "url", "primary"}
    for secret in (
        "REG-123",
        "attendee@example",
        "person-1",
        "object_key",
        "user_id",
        "staff_notes",
        "helper",
    ):
        assert secret not in response.text
    assert "no-cache" in response.headers["cache-control"]
    assert (
        client.get(
            "/api/v1/events/2028/creators",
            headers={"If-None-Match": response.headers["etag"]},
        ).status_code
        == 304
    )
    html = client.get("/gallery/2028")
    assert "Zulu &lt;creator&gt;" in html.text and "(Primary)" in html.text
    assert html.text.index("Alpha") < html.text.index("Zulu")
    assert "/gallery/2028" in client.get("/").text
    assert client.get("/api/v1/events/1900/creators").json() == []
    assert client.post("/api/v1/events/2028/creators").status_code == 405
    assert (
        "PublicCreator" in client.get("/openapi.json").json()["components"]["schemas"]
    )


@pytest.mark.parametrize(
    "change",
    [
        "withdrawn",
        "rejected",
        "hidden",
        "name",
        "image",
        "inactive_image",
        "expired",
        "cleanup",
    ],
)
def test_publication_predicate(gallery_setup, change):
    client, engine, key, public_id, _ = gallery_setup
    with Session(engine) as db:
        record, profile, event = (
            db.get(CreatorApplication, key),
            db.get(CreatorProfile, key),
            db.get(Event, 1),
        )
        if change == "withdrawn":
            record.withdrawn_at = datetime.now(UTC)
        elif change == "rejected":
            record.status, record.outcome_reason = "NOT_APPROVED", "reason"
        elif change == "hidden":
            profile.publicly_hidden = True
        elif change == "name":
            profile.channel_name = ""
        elif change == "image":
            profile.image_id = None
        elif change == "inactive_image":
            db.get(ProfileImage, profile.image_id).state = "DELETE"
        elif change == "expired":
            event.data_delete_at = datetime.now(UTC) - timedelta(seconds=1)
        else:
            event.cleanup_started_at = datetime.now(UTC)
        db.commit()
    assert public_id not in client.get("/api/v1/events/2028/creators").text
    assert (
        client.get(f"/api/v1/events/2028/creators/{public_id}/image").status_code == 404
    )


def test_admin_visibility_auth_csrf_scope_and_reversible(gallery_setup):
    client, engine, key, public_id, _ = gallery_setup
    url = f"/admin/creators/{key}/visibility"
    assert client.post(url, data={"hidden": "yes"}).status_code == 401
    m2.login(client, 1)
    assert client.post(url, data={"hidden": "yes"}).status_code == 403
    m2.login(client, 2)
    assert client.post(url, data={"hidden": "yes"}).status_code == 403
    path = f"/creators/{key}?administrative=true"
    with Session(engine) as db:
        channel = db.scalar(
            select(CreatorChannel.id).where(
                CreatorChannel.application_id == key,
                CreatorChannel.is_primary.is_(True),
            )
        )
        foreign = db.scalar(
            select(CreatorChannel.id).where(CreatorChannel.application_id != key)
        )
    assert (
        client.post(
            url,
            data={
                "csrf_token": m2.csrf(client, path),
                "hidden": "yes",
                "channel_id": foreign,
            },
        ).status_code
        == 404
    )
    assert (
        client.post(
            url,
            data={
                "csrf_token": m2.csrf(client, path),
                "hidden": "yes",
                "channel_id": channel,
            },
            follow_redirects=False,
        ).status_code
        == 303
    )
    row = next(
        r
        for r in client.get("/api/v1/events/2028/creators").json()
        if r["id"] == public_id
    )
    assert len(row["channels"]) == 1 and row["channels"][0]["platform"] == "Mastodon"
    # Hiding a primary public link doesn't change badge/print data or workflow.
    for hidden in ("yes", "no"):
        assert (
            client.post(
                url,
                data={"csrf_token": m2.csrf(client, path), "hidden": hidden},
                follow_redirects=False,
            ).status_code
            == 303
        )
        assert (public_id in client.get("/api/v1/events/2028/creators").text) == (
            hidden == "no"
        )
    with Session(engine) as db:
        assert db.get(CreatorApplication, key).status == "APPROVED"
        assert (
            db.scalar(
                select(func.count())
                .select_from(BusinessAudit)
                .where(BusinessAudit.action == "public_visibility_changed")
            )
            == 3
        )
        db.execute(delete(LocalRoleAssignment).where(LocalRoleAssignment.user_id == 2))
        db.commit()
    assert client.post(url, data={"hidden": "yes"}).status_code == 403


def test_public_image_rechecks_visibility_and_storage_failure(gallery_setup):
    client, engine, key, public_id, store = gallery_setup
    url = f"/api/v1/events/2028/creators/{public_id}/image"
    response = client.get(url)
    assert (
        response.status_code == 200 and response.headers["content-type"] == "image/png"
    )
    assert response.headers["x-content-type-options"] == "nosniff"
    assert (
        client.get(
            url, headers={"If-None-Match": "W/" + response.headers["etag"]}
        ).status_code
        == 304
    )
    store.failure = True
    assert client.get(url).status_code == 503
    store.failure = False

    def hide():
        with Session(engine) as db:
            service.set_visibility(db, 2, key, True)

    store.hook = hide
    assert client.get(url).status_code == 404
    assert (
        client.get(url, headers={"If-None-Match": response.headers["etag"]}).status_code
        == 404
    )


def expire(db):
    event = db.get(Event, 1)
    event.data_delete_at = datetime.now(UTC) - timedelta(seconds=1)
    db.commit()


def test_full_cleanup_failure_retry_and_only_bans_retained(application_engine):
    with Session(application_engine, expire_on_commit=False) as db:
        key, _, object_key = published(db)
        helper = m3.register(db, key)
        m3.decide(db, helper)
        db.add(
            BannedChannel(
                platform="Twitch",
                normalized_account="banned",
                private_reason="private",
                created_at=datetime.now(UTC),
                updated_at=datetime.now(UTC),
            )
        )
        db.commit()
        store = m3.StoreFake(db)
        store.objects[object_key] = m3.png()
        expire(db)
        assert claim(db, datetime.now(UTC)) is None
        store.fail_delete = True
        assert cleanup_due_events(db, store) == 1
        event = db.get(Event, 1)
        assert event.cleanup_failed and event.cleanup_started_at
        assert db.get(CreatorApplication, key) and db.get(HelperRegistration, helper)
        with pytest.raises(HTTPException) as exc:
            workflow.active_event(db, lock=True)
        assert exc.value.status_code == 410
        store.fail_delete = False
        assert cleanup_due_events(db, store) == 0
        assert store.objects == {}
        assert cleanup_event(db, store, 1)
        for table in Base.metadata.sorted_tables:
            assert db.scalar(select(func.count()).select_from(table)) == (
                1 if table.name == "banned_channels" else 0
            ), table.name


@pytest.mark.parametrize("kind", ["upload", "notification"])
def test_cleanup_waits_for_inflight_operations(application_engine, kind):
    with Session(application_engine, expire_on_commit=False) as db:
        published(db)
        if kind == "upload":
            db.add(
                ProfileImage(
                    event_id=1,
                    object_key="staged",
                    state="STAGED",
                    delete_after=datetime.now(UTC) + timedelta(hours=1),
                )
            )
        else:
            notice = db.scalar(select(NotificationOutbox))
            notice.state = "SENDING"
            notice.claimed_until = datetime.now(UTC) + timedelta(minutes=5)
        db.commit()
        expire(db)
        store = m3.StoreFake(db)
        assert not cleanup_event(db, store, 1)
        if kind == "upload":
            db.scalar(
                select(ProfileImage).where(ProfileImage.object_key == "staged")
            ).delete_after = datetime.now(UTC) - timedelta(seconds=1)
        else:
            db.scalar(select(NotificationOutbox)).claimed_until = datetime.now(
                UTC
            ) - timedelta(seconds=1)
        db.commit()
        assert cleanup_event(db, store, 1)


def test_cleanup_preserves_identity_needed_by_other_event(
    application_engine, monkeypatch
):
    with Session(application_engine, expire_on_commit=False) as db:
        published(db)
        now = datetime.now(UTC)
        db.add(
            Event(
                id=2,
                year=2029,
                name="Next",
                starts_at=now,
                ends_at=now,
                application_open_at=now - timedelta(days=1),
                application_close_at=now + timedelta(days=1),
                badge_change_deadline_at=now + timedelta(days=1),
                badge_print_at=now,
            )
        )
        db.commit()
        monkeypatch.setattr(workflow.settings, "active_event_id", 2)
        second = asyncio.run(
            workflow.submit(
                db,
                1,
                m2.input_data(),
                m2.RegistrationFake(),
                m2.IdentityFake(),
                **m2.picture_input(),
            )
        )
        expire(db)
        assert cleanup_event(db, m3.StoreFake(db), 1)
        assert db.get(LocalUser, 1) is not None
        assert db.get(CreatorApplication, second).event_id == 2
        event = db.get(Event, 2)
        assert workflow.utc(event.data_delete_at) == workflow.utc(
            event.ends_at
        ) + timedelta(days=30)
        event.data_delete_at = now - timedelta(seconds=1)
        db.commit()
        assert cleanup_event(db, m3.StoreFake(db), 2)
        assert db.get(LocalUser, 1, populate_existing=True) is None


def test_deleted_recreated_identity_cannot_reuse_signed_cookie(browser):
    client, _, engine = browser
    old_cookie = client.cookies.get("creator_session")
    with Session(engine) as db:
        db.execute(delete(ExternalIdentity).where(ExternalIdentity.user_id == 1))
        db.execute(delete(LocalUser).where(LocalUser.id == 1))
        db.add(LocalUser(id=1))
        db.commit()
    client.cookies.clear()
    client.cookies.set("creator_session", old_cookie)
    assert client.get("/auth/me").status_code == 401


def test_readiness_is_database_only_and_liveness_independent(browser):
    client, _, _ = browser
    assert client.get("/ready").status_code == 200
    original = app.dependency_overrides[get_db]

    class BrokenDB:
        def execute(self, _):
            raise OperationalError("private DSN", {}, Exception("secret"))

    app.dependency_overrides[get_db] = BrokenDB
    try:
        response = client.get("/ready")
        assert response.status_code == 503 and "secret" not in response.text
        assert client.get("/health").status_code == 200
    finally:
        app.dependency_overrides[get_db] = original


def test_postgres_simultaneous_cleanup_is_idempotent(application_engine):
    if application_engine.dialect.name != "postgresql":
        pytest.skip("PostgreSQL cleanup locking")
    with Session(application_engine, expire_on_commit=False) as db:
        published(db)
        expire(db)
    barrier = Barrier(2)

    def worker(_):
        with Session(application_engine, expire_on_commit=False) as db:
            barrier.wait(timeout=10)
            return cleanup_event(db, m3.StoreFake(db), 1)

    with ThreadPoolExecutor(max_workers=2) as pool:
        assert all(pool.map(worker, range(2)))
    with Session(application_engine) as db:
        assert db.get(Event, 1) is None


def test_helm_rendering():
    helm = os.environ.get("HELM_BINARY")
    if not helm:
        pytest.skip("Set HELM_BINARY for deployment manifest verification")
    args = [
        helm,
        "template",
        "test",
        "charts/creators",
        "--set",
        "image.repository=creators",
        "--set",
        "image.tag=test",
        "--set",
        "workers.cleanup.enabled=true",
        "--set",
        "workers.notifications.enabled=true",
        "--set",
        "ingress.enabled=true",
        "--set",
        "ingress.host=creators.example.test",
    ]
    result = subprocess.run(args, capture_output=True, text=True, check=True)
    manifest = result.stdout
    assert manifest.count("kind: Job") == 1
    assert manifest.count("kind: CronJob") == 2
    assert "pre-install,pre-upgrade" in manifest
    assert "alembic, upgrade, head" in manifest and "downgrade" not in manifest
    assert "path: /ready" in manifest and "path: /health" in manifest
    assert (
        "runAsNonRoot: true" in manifest
        and "automountServiceAccountToken: false" in manifest
    )
    assert "app.events.cleanup" in manifest and "app.notifications.worker" in manifest
    assert "DATABASE_URL:" not in manifest


def test_visibility_survives_admin_channel_correction(application_engine):
    with Session(application_engine, expire_on_commit=False) as db:
        application_id, public_id, _ = published(db)
        channel_id = db.scalar(
            select(CreatorChannel.id).where(
                CreatorChannel.application_id == application_id
            )
        )
        service.set_visibility(db, 2, application_id, True, channel_id=channel_id)
        application = db.get(CreatorApplication, application_id)
        workflow.edit(
            db,
            2,
            application_id,
            application.version,
            m2.input_data("@creator1"),
            administrative=True,
            reason="Correct content",
        )
        assert service.public_records(db, 2028)[0][0].channels == []
        assert db.get(CreatorProfile, application_id).public_id == public_id


def test_event_configuration_and_retention_dashboard(browser):
    from app.events.manage import EventConfiguration, configure

    client, _, engine = browser
    fields = {
        "id": 2,
        "year": 2029,
        "name": "Next event",
        "starts_at": "2029-09-01T10:00:00+02:00",
        "ends_at": "2029-09-04T10:00:00+02:00",
        "application_open_at": "2029-01-01T00:00:00+01:00",
        "application_close_at": "2029-06-01T00:00:00+02:00",
        "badge_change_deadline_at": "2029-08-01T00:00:00+02:00",
        "badge_print_at": "2029-08-02T00:00:00+02:00",
    }
    configuration = EventConfiguration.model_validate(fields)
    assert configuration.starts_at.hour == 8
    assert configuration.data_delete_at == configuration.ends_at + timedelta(days=30)
    with pytest.raises(ValueError):
        EventConfiguration.model_validate(
            {**fields, "starts_at": "2029-09-01T10:00:00"}
        )
    with Session(engine) as db, db.begin():
        configure(db, configuration, "Next event")
    assert client.get("/admin/retention").status_code == 403
    m2.login(client, 2)
    assert "Next event" in client.get("/admin/retention").text
    with Session(engine) as db, db.begin():
        db.get(Event, 2).cleanup_started_at = datetime.now(UTC)
    with (
        Session(engine) as db,
        db.begin(),
        pytest.raises(ValueError, match="Cleanup has started"),
    ):
        configure(db, configuration, "Must refuse")


def test_cleanup_database_failure_is_retryable(application_engine):
    from sqlalchemy import event as sql_event

    with Session(application_engine, expire_on_commit=False) as db:
        published(db)
        expire(db)

        def fail_delete(session, *_):
            if any(isinstance(row, Event) for row in session.deleted):
                raise OperationalError("DELETE", {}, Exception("failure"))

        sql_event.listen(db, "before_flush", fail_delete)
        assert cleanup_due_events(db, m3.StoreFake(db)) == 1
        sql_event.remove(db, "before_flush", fail_delete)
        assert db.get(Event, 1) is not None
        assert cleanup_due_events(db, m3.StoreFake(db)) == 0


def test_representative_local_end_to_end(browser, provider):
    import re
    from io import BytesIO

    from openpyxl import load_workbook

    client, _, engine = browser
    provider.claims["sub"] = "person-1"
    client.cookies.clear()
    auth_tests.begin(client, provider)
    callback = auth_tests.finish(client, provider)
    assert callback.status_code == 303 and callback.headers["location"] == "/"
    assert m2.submit_form(client).status_code == 303
    with Session(engine) as db:
        application_id = db.scalar(select(CreatorApplication.id))
    m2.login(client, 2)
    review_path = f"/admin/applications/{application_id}"
    for version, status in ((1, "ON_REVIEW"), (2, "APPROVED")):
        assert (
            client.post(
                review_path,
                data={
                    "csrf_token": m2.csrf(client, review_path),
                    "version": version,
                    "status": status,
                },
                follow_redirects=False,
            ).status_code
            == 303
        )
    m2.login(client, 1)
    path = f"/creators/{application_id}"
    assert (
        client.post(
            path + "/profile",
            data={
                "csrf_token": m2.csrf(client, path),
                "version": 3,
                "channel_name": "Creator",
                "platform": "Twitch",
                "account": "creator",
                "primary": "0",
            },
            follow_redirects=False,
        ).status_code
        == 303
    )

    class MemoryStore:
        def __init__(self):
            self.objects = {}

        def put(self, key, data):
            self.objects[key] = data

        def get(self, key):
            return self.objects[key]

        def delete(self, key):
            self.objects.pop(key, None)

    store = MemoryStore()
    app.dependency_overrides[images.get_image_store] = lambda: store
    assert (
        client.post(
            path + "/picture",
            data={"csrf_token": m2.csrf(client, path), "version": 4},
            files={"picture": ("../../unsafe.png", m3.png(), "image/png")},
            follow_redirects=False,
        ).status_code
        == 303
    )
    response = client.post(
        path + "/invitations", data={"csrf_token": m2.csrf(client, path)}
    )
    assert response.status_code == 200
    secret = re.search(r"/helpers/redeem#([A-Za-z0-9_-]+)", response.text)[1]
    m2.login(client, 3)
    assert (
        client.post(
            "/helpers/redeem",
            data={
                "csrf_token": m2.csrf(client, "/helpers/redeem"),
                "invitation": secret,
            },
            follow_redirects=False,
        ).status_code
        == 303
    )
    with Session(engine) as db:
        helper_id = db.scalar(select(HelperRegistration.id))
    m2.login(client, 1)
    assert (
        client.post(
            f"/helpers/{helper_id}/decision",
            data={
                "csrf_token": m2.csrf(client, path),
                "version": 1,
                "status": "CONFIRMED",
            },
            follow_redirects=False,
        ).status_code
        == 303
    )
    assert len(client.get("/api/v1/events/2028/creators").json()) == 1
    m2.login(client, 2)
    response = client.post(
        "/admin/events/1/export",
        data={
            "csrf_token": m2.csrf(client, "/admin/events/1/operations"),
            "mode": "print",
        },
    )
    assert response.status_code == 200
    assert load_workbook(BytesIO(response.content))["Print"].max_row == 3
    m2.login(client, 3)
    response = client.post(
        "/staff/events/1/badges",
        data={
            "csrf_token": m2.csrf(client, "/staff/events/1/badges"),
            "reg_id": "REG-123",
        },
    )
    assert "Badge 1" in response.text and "Badge 2" in response.text
    with Session(engine) as db:
        badge_id = db.scalar(select(Badge.id).where(Badge.helper_id == helper_id))
    assert (
        client.post(
            f"/staff/events/1/badges/{badge_id}",
            data={
                "csrf_token": m2.csrf(client, "/staff/events/1/badges"),
                "action": "pickup",
            },
            follow_redirects=False,
        ).status_code
        == 303
    )
    with Session(engine, expire_on_commit=False) as db:
        m3.withdraw(db, application_id)
        assert service.public_records(db, 2028) == []
        expire(db)
        assert cleanup_event(db, store, 1)
    assert client.get("/api/v1/events/2028/creators").json() == []
    assert client.get("/auth/me").status_code == 401
    assert store.objects == {}


def test_png_compressed_metadata_bomb_is_rejected():
    from io import BytesIO

    from PIL import Image, PngImagePlugin

    metadata = PngImagePlugin.PngInfo()
    metadata.add_text("hostile", "x" * (PngImagePlugin.MAX_TEXT_CHUNK + 1), zip=True)
    output = BytesIO()
    Image.new("RGB", (600, 600)).save(output, format="PNG", pnginfo=metadata)
    with pytest.raises(HTTPException) as exc:
        images.normalize_png(output.getvalue())
    assert exc.value.status_code == 422


def test_database_errors_do_not_log_private_values(browser, caplog):
    client, _, _ = browser
    original = app.dependency_overrides[get_db]

    class BrokenDB:
        def scalars(self, _):
            raise OperationalError(
                "private SQL",
                {"email": "private@example.test"},
                Exception("private row"),
            )

    app.dependency_overrides[get_db] = BrokenDB
    try:
        assert client.get("/").status_code == 503
        assert "Database operation failed" in caplog.text
        assert "private" not in caplog.text
    finally:
        app.dependency_overrides[get_db] = original
