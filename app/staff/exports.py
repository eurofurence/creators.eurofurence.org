"""Private, on-demand operational downloads with explicit data projections."""

import json
import re
from datetime import UTC, datetime
from io import BytesIO

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from sqlalchemy import or_, select

from app.applications.models import (
    Badge,
    BadgeCounter,
    BusinessAudit,
    ConventionVideo,
    CreatorApplication,
    CreatorChannel,
    LocalRoleAssignment,
    NotificationOutbox,
)
from app.applications.security import require_admin
from app.applications.service import registration_lookup_for_user
from app.applications.workflow import utc
from app.creators.images import ImageStorageUnavailable, normalize_png
from app.creators.models import CreatorProfile, ProfileImage
from app.creators.policy import creator_active, helper_active
from app.helpers.models import HelperInvitation, HelperRegistration
from app.moderation.service import warnings
from app.registration.eligibility import Eligibility, check_eligibility
from app.staff.service import event_record

PRINT_COLUMNS = ("NUMBER", "NAME", "ICON", "LINK", "PICTURE")


def picture_filename(channel, application_id):
    account = channel.normalized_account.removeprefix("@")
    safe = re.sub(r"[^A-Za-z0-9_.-]", "_", account).strip("._-")[:100] or "creator"
    # Disambiguate identical handles on different platforms/creators and reserved OS names.
    return f"{safe}-{application_id}.png"


def project(record, fields):
    return {field: getattr(record, field) for field in fields.split()}


def snapshot(db, event_id):
    event = event_record(db, event_id)
    applications = list(
        db.scalars(
            select(CreatorApplication)
            .where(CreatorApplication.event_id == event_id)
            .order_by(CreatorApplication.id)
        )
    )
    ids = [a.id for a in applications]
    channels = list(
        db.scalars(
            select(CreatorChannel)
            .where(CreatorChannel.application_id.in_(ids))
            .order_by(CreatorChannel.id)
        )
    )
    profiles = {
        p.application_id: p
        for p in db.scalars(
            select(CreatorProfile).where(CreatorProfile.application_id.in_(ids))
        )
    }
    images = {
        i.id: i
        for i in db.scalars(
            select(ProfileImage).where(ProfileImage.event_id == event_id)
        )
    }
    helpers = list(
        db.scalars(
            select(HelperRegistration)
            .where(HelperRegistration.event_id == event_id)
            .order_by(HelperRegistration.id)
        )
    )
    badges = list(
        db.scalars(
            select(Badge).where(Badge.event_id == event_id).order_by(Badge.badge_number)
        )
    )
    primary = {c.application_id: c for c in channels if c.is_primary}
    app_map = {a.id: a for a in applications}
    helper_map = {h.id: h for h in helpers}
    creator_badges = {b.application_id: b for b in badges if b.application_id}
    helper_badges = {b.helper_id: b for b in badges if b.helper_id}
    candidates = []
    owners = [
        (a, None, creator_badges.get(a.id)) for a in applications if creator_active(a)
    ]
    owners += [
        (app_map[h.application_id], h, helper_badges.get(h.id))
        for h in helpers
        if helper_active(h, app_map[h.application_id])
    ]
    for application, helper, badge in owners:
        profile, channel = profiles.get(application.id), primary.get(application.id)
        image = images.get(profile.image_id) if profile else None
        user_id = helper.user_id if helper else application.user_id
        candidates.append(
            {
                "badge_id": badge.id if badge else None,
                "number": badge.badge_number if badge else None,
                "application_id": application.id,
                "helper_id": helper.id if helper else None,
                "user_id": user_id,
                "version": application.version,
                "helper_version": helper.version if helper else None,
                "name": profile.channel_name if profile else "",
                "platform": channel.platform if channel else "",
                "link": channel.canonical_url if channel else "",
                "filename": picture_filename(channel, application.id)
                if channel
                else "",
                "image_id": image.id if image and image.state == "ACTIVE" else None,
                "object_key": image.object_key
                if image and image.state == "ACTIVE"
                else None,
                "lookup": registration_lookup_for_user(db, user_id, event_id=event_id),
            }
        )
    candidates.sort(
        key=lambda c: (
            c["number"] is None,
            c["number"] or 0,
            c["application_id"],
            c["helper_id"] or 0,
        )
    )
    data = {
        "Event": [
            project(
                event,
                "id year name starts_at ends_at application_open_at application_close_at badge_change_deadline_at badge_print_at data_delete_at helper_limit",
            )
        ],
        "Applications": [
            project(
                a,
                "id event_id user_id status version withdrawn_at livestream shorts vlogs reg_id nickname email eligibility_checked_at outcome_reason staff_notes created_at updated_at",
            )
            for a in applications
        ],
        "Channels": [
            project(
                c,
                "id application_id platform original_representation normalized_account canonical_url is_primary publicly_hidden",
            )
            for c in channels
        ],
        "Videos": [
            project(v, "application_id position url")
            for v in db.scalars(
                select(ConventionVideo)
                .where(ConventionVideo.application_id.in_(ids))
                .order_by(ConventionVideo.application_id, ConventionVideo.position)
            )
        ],
        "Profiles": [
            {
                **project(
                    p, "application_id channel_name image_id public_id publicly_hidden"
                ),
                "picture_filename": picture_filename(
                    primary[p.application_id], p.application_id
                )
                if p.application_id in primary
                else None,
            }
            for p in profiles.values()
        ],
        "Images": [
            project(i, "id event_id object_key state delete_after deletion_failed")
            for i in images.values()
        ],
        "Invitations": [
            {
                **project(
                    i,
                    "id application_id event_id created_at expires_at revoked_at consumed_at consumed_by",
                ),
                "state": "USED"
                if i.consumed_at
                else "REVOKED"
                if i.revoked_at
                else "EXPIRED"
                if utc(i.expires_at) <= datetime.now(UTC)
                else "UNUSED",
            }
            for i in db.scalars(
                select(HelperInvitation)
                .where(HelperInvitation.event_id == event_id)
                .order_by(HelperInvitation.id)
            )
        ],
        "Helpers": [
            {
                **project(
                    h,
                    "id application_id event_id creator_user_id user_id invitation_id status version withdrawn_at reg_id nickname email eligibility_checked_at created_at",
                ),
                "active": helper_active(h, app_map[h.application_id]),
            }
            for h in helpers
        ],
        "Badges": [
            {
                **project(
                    b,
                    "id event_id application_id helper_id badge_number assigned_at picked_up_at picked_up_by",
                ),
                "active": (
                    helper_active(
                        helper_map[b.helper_id],
                        app_map[helper_map[b.helper_id].application_id],
                    )
                    if b.helper_id
                    else creator_active(app_map[b.application_id])
                ),
            }
            for b in badges
        ],
        "Counters": [
            project(c, "event_id next_number")
            for c in db.scalars(
                select(BadgeCounter).where(BadgeCounter.event_id == event_id)
            )
        ],
        "Roles": [
            project(r, "user_id event_id role")
            for r in db.scalars(
                select(LocalRoleAssignment)
                .where(
                    or_(
                        LocalRoleAssignment.event_id == event_id,
                        LocalRoleAssignment.role == "ADMIN",
                    )
                )
                .order_by(LocalRoleAssignment.id)
            )
        ],
        "Notifications": [
            project(
                n,
                "id event_id application_id application_version recipient_id notification_type state attempts next_attempt_at claimed_until sent_at last_error created_at",
            )
            for n in db.scalars(
                select(NotificationOutbox)
                .where(NotificationOutbox.event_id == event_id)
                .order_by(NotificationOutbox.id)
            )
        ],
        "Audit": [
            project(
                a,
                "id event_id actor_id entity entity_id action reason changes occurred_at",
            )
            for a in db.scalars(
                select(BusinessAudit)
                .where(
                    or_(
                        BusinessAudit.event_id == event_id,
                        BusinessAudit.event_id.is_(None),
                    )
                )
                .order_by(BusinessAudit.id)
            )
        ],
        "Ban warnings": [
            {
                "application_id": a.id,
                **project(
                    b,
                    "id platform normalized_account original_reference private_reason active created_at updated_at",
                ),
            }
            for a in applications
            for b in warnings(db, a.id)
        ],
    }
    return candidates, data


def problem(candidate, state):
    return {
        "badge_id": candidate["badge_id"],
        "number": candidate["number"],
        "application_id": candidate["application_id"],
        "helper_id": candidate["helper_id"],
        "problem": state,
    }


async def validate_print(candidates, registration, store):
    problems = []
    checked_images = {}
    for row in candidates:
        eligibility = (
            await check_eligibility(registration, row["lookup"])
            if row["lookup"]
            else Eligibility.UNAVAILABLE
        )
        if eligibility is not Eligibility.ELIGIBLE:
            problems.append(problem(row, eligibility.value))
        missing = [
            key
            for key in ("number", "name", "platform", "link", "image_id")
            if not row[key]
        ]
        if missing:
            problems.append(problem(row, "Missing " + ", ".join(missing)))
        if row["object_key"]:
            key = row["object_key"]
            if key not in checked_images:
                try:
                    normalize_png(store.get(key))
                    checked_images[key] = True
                except ImageStorageUnavailable, HTTPException:
                    checked_images[key] = False
            if not checked_images[key]:
                problems.append(problem(row, "Profile image unavailable or invalid"))
    return problems


def write_sheet(workbook, name, headers, rows):
    sheet = workbook.create_sheet(name)
    for values in (headers, *rows):
        sheet.append([None] * len(values))
        for index, value in enumerate(values, 1):
            cell = sheet.cell(sheet.max_row, index)
            if isinstance(value, dict):
                value = json.dumps(value, ensure_ascii=False, sort_keys=True)
            if isinstance(value, datetime):
                value = utc(value).replace(tzinfo=None)
                cell.number_format = 'yyyy-mm-dd hh:mm:ss" UTC"'
            if isinstance(value, str):
                value = ILLEGAL_CHARACTERS_RE.sub("", value)
                if len(value) > 32767:
                    raise HTTPException(
                        422,
                        "A retained text field exceeds Excel's cell limit; export cannot preserve it completely",
                    )
                cell.value = value
                cell.data_type = (
                    "s"  # Literal text, including =, +, -, @, tabs and newlines.
                )
            else:
                cell.value = value
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243746")
        sheet.column_dimensions[cell.column_letter].width = min(
            50, max(18, len(str(cell.value)) + 3)
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def workbook_bytes(data, candidates=None):
    workbook = Workbook()
    workbook.remove(workbook.active)
    if candidates is not None:
        write_sheet(
            workbook,
            "Print",
            PRINT_COLUMNS,
            [
                [r[k] for k in ("number", "name", "platform", "link", "filename")]
                for r in candidates
            ],
        )
    else:
        write_sheet(
            workbook,
            "Export status",
            ("Status",),
            [("Operations only. Not validated for printing.",)],
        )
    for name, records in data.items():
        headers = tuple(records[0]) if records else ("No retained records",)
        write_sheet(workbook, name, headers, [[r[h] for h in headers] for r in records])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


async def export(db, actor_id, event_id, registration, store, *, print_ready=True):
    require_admin(db, actor_id)
    candidates, data = snapshot(db, event_id)
    event_context = data["Event"]
    db.rollback()
    problems = (
        await validate_print(candidates, registration, store) if print_ready else []
    )
    with db.begin():
        require_admin(db, actor_id, lock=True)
        event_record(db, event_id, lock=True)
        current, data = snapshot(db, event_id)
        if print_ready and (current != candidates or data["Event"] != event_context):
            raise HTTPException(
                409,
                "Print data changed during verification. Generate the export again.",
            )
        if problems:
            raise HTTPException(
                409,
                {
                    "message": "Print-ready export blocked. Resolve these problems and retry. Operations-only export remains available.",
                    "problems": problems,
                },
            )
        db.add(
            BusinessAudit(
                actor_id=actor_id,
                event_id=event_id,
                entity="event",
                entity_id=event_id,
                action="print_exported" if print_ready else "operations_exported",
                changes={"print_rows": len(current) if print_ready else 0},
            )
        )
        db.flush()
        # Include this export's audit record as part of the operational fallback.
        _, data = snapshot(db, event_id)
        return workbook_bytes(data, current if print_ready else None)


def download_picture(db, actor_id, application_id, store):
    require_admin(db, actor_id)
    application = db.get(CreatorApplication, application_id)
    profile = db.get(CreatorProfile, application_id)
    channel = db.scalar(
        select(CreatorChannel).where(
            CreatorChannel.application_id == application_id,
            CreatorChannel.is_primary.is_(True),
        )
    )
    image = (
        db.get(ProfileImage, profile.image_id) if profile and profile.image_id else None
    )
    if not application or not image or image.state != "ACTIVE" or not channel:
        raise HTTPException(404, "Profile picture or primary channel is unavailable")
    key, image_id, event_id = image.object_key, image.id, application.event_id
    filename = picture_filename(channel, application_id)
    db.rollback()
    try:
        data = store.get(key)
        normalize_png(data)
    except ImageStorageUnavailable, HTTPException:
        raise HTTPException(
            503, "Profile picture is unavailable. Retry or correct the stored picture."
        ) from None
    with db.begin():
        require_admin(db, actor_id, lock=True)
        event_record(db, event_id, lock=True)
        profile = db.get(CreatorProfile, application_id, populate_existing=True)
        current_channel = db.scalar(
            select(CreatorChannel)
            .where(
                CreatorChannel.application_id == application_id,
                CreatorChannel.is_primary.is_(True),
            )
            .execution_options(populate_existing=True)
        )
        current_image = db.get(ProfileImage, image_id, populate_existing=True)
        if (
            not profile
            or profile.image_id != image_id
            or not current_channel
            or picture_filename(current_channel, application_id) != filename
            or not current_image
            or current_image.state != "ACTIVE"
        ):
            raise HTTPException(409, "Profile picture changed. Retry the download.")
        db.add(
            BusinessAudit(
                actor_id=actor_id,
                event_id=event_id,
                entity="profile",
                entity_id=application_id,
                action="picture_downloaded",
                changes={},
            )
        )
    return filename, data
